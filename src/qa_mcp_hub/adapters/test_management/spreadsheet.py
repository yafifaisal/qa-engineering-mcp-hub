"""CSV and XLSX test-case import and export.

Purpose: read and write local workbooks without evaluating formulas.
Milestone: 09
Status: IMPLEMENTED
Ownership: adapters.test_management.spreadsheet
Dependencies: Canonical cases and openpyxl.
Expected tests: encoding, Unicode, roundtrip, and formula neutralization.
"""

import csv
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]

from qa_mcp_hub.adapters.test_management.base import (
    CanonicalCase,
    FieldMapping,
    PreservedField,
    content_digest,
    import_rows,
    join_steps,
    neutralize_formula,
)
from qa_mcp_hub.core.errors import HubError
from qa_mcp_hub.security.path_allowlist import resolve_within

_HEADERS = (
    "ID",
    "Version",
    "Scenario",
    "Priority",
    "Test Case",
    "Precondition",
    "Steps",
    "Expected Result",
    "Result",
    "Notes",
    "Evidence",
    "Automation Suite",
)


def read_csv(path: Path, workspace: Path) -> list[dict[str, str]]:
    """Read a UTF-8 CSV. UTF-16 and other encodings are rejected."""
    resolved = resolve_within(path, [workspace])
    raw = resolved.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise HubError("encoding_rejected", "CSV encoding must be UTF-8")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HubError("encoding_rejected", "CSV encoding must be UTF-8") from exc
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise HubError("missing_field", "CSV header is missing")
    rows: list[dict[str, str]] = []
    for raw_row in reader:
        row: dict[str, str] = {}
        for key, value in raw_row.items():
            if key is None:
                continue
            row[key] = "" if value is None else value
        rows.append(row)
    return rows


def read_xlsx(path: Path, workspace: Path) -> list[dict[str, str]]:
    """Read the first worksheet. Formulas stay text and are not evaluated."""
    resolved = resolve_within(path, [workspace])
    workbook = load_workbook(resolved, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet is None:
            raise HubError("missing_field", "workbook has no worksheet")
        grid = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not grid:
        raise HubError("missing_field", "workbook header is missing")
    headers = [_cell_text(value) for value in grid[0]]
    if any(header == "" for header in headers):
        raise HubError("missing_field", "workbook header is missing")
    rows: list[dict[str, str]] = []
    for values in grid[1:]:
        row = {
            headers[index]: _cell_text(values[index] if index < len(values) else None)
            for index in range(len(headers))
        }
        if any(row.values()):
            rows.append(row)
    return rows


def import_csv(
    path: Path, workspace: Path, mapping: FieldMapping
) -> tuple[CanonicalCase, ...]:
    """Import a CSV through the reviewed column map."""
    return _import_table(read_csv(path, workspace), mapping)


def import_xlsx(
    path: Path, workspace: Path, mapping: FieldMapping
) -> tuple[CanonicalCase, ...]:
    """Import a workbook through the reviewed column map."""
    return _import_table(read_xlsx(path, workspace), mapping)


def export_csv(
    cases: tuple[CanonicalCase, ...], path: Path, workspace: Path
) -> tuple[str, tuple[str, ...], str, bool]:
    """Write a UTF-8 CSV. An identical file is left unchanged."""
    resolved = resolve_within(path, [workspace])
    text, neutralized = _csv_text(cases)
    digest = content_digest(text)
    if (
        resolved.is_file()
        and content_digest(resolved.read_text(encoding="utf-8")) == digest
    ):
        return text, neutralized, digest, False
    resolved.parent.mkdir(parents=True, exist_ok=True)
    resolved.write_text(text, encoding="utf-8", newline="\n")
    return text, neutralized, digest, True


def export_xlsx(
    cases: tuple[CanonicalCase, ...], path: Path, workspace: Path
) -> tuple[str, tuple[str, ...], str, bool]:
    """Write a workbook whose cells cannot start a formula."""
    resolved = resolve_within(path, [workspace])
    text, neutralized = _csv_text(cases)
    digest = content_digest(text)
    if resolved.is_file() and _xlsx_digest(resolved) == digest:
        return text, neutralized, digest, False
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise HubError("mapping_rejected", "workbook could not be created")
    sheet.title = "Cases"
    for row in csv.reader(io.StringIO(text)):
        sheet.append(row)
    resolved.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(resolved)
    workbook.close()
    sidecar = resolved.with_suffix(resolved.suffix + ".digest")
    sidecar.write_text(digest, encoding="utf-8")
    return text, neutralized, digest, True


def _import_table(
    rows: list[dict[str, str]], mapping: FieldMapping
) -> tuple[CanonicalCase, ...]:
    known = {
        name
        for name in (
            mapping.columns.source_id,
            mapping.columns.test_case,
            mapping.columns.steps,
            mapping.columns.expected_result,
            mapping.columns.priority,
            mapping.columns.scenario,
            mapping.columns.precondition,
            mapping.columns.result,
            mapping.columns.notes,
            mapping.columns.evidence,
            mapping.columns.version,
            mapping.columns.automation_suite_id,
        )
        if name
    }
    cleaned: list[dict[str, str]] = []
    extras: list[tuple[PreservedField, ...]] = []
    for row in rows:
        preserved = tuple(
            PreservedField(name=key, value=value)
            for key, value in row.items()
            if key not in known and value
        )
        cleaned.append(row)
        extras.append(preserved)
    return import_rows(cleaned, mapping, extras=extras)


def _csv_text(cases: tuple[CanonicalCase, ...]) -> tuple[str, tuple[str, ...]]:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n", quoting=csv.QUOTE_ALL)
    writer.writerow(_HEADERS)
    neutralized: list[str] = []
    for case in cases:
        raw = [
            case.source_id,
            case.version,
            case.scenario,
            case.priority,
            case.test_case,
            case.precondition,
            join_steps(case.steps),
            case.expected_result,
            case.result,
            case.notes,
            case.evidence,
            case.automation_suite_id,
        ]
        safe: list[str] = []
        for value in raw:
            text, changed = neutralize_formula(value)
            if changed:
                neutralized.append(value)
            safe.append(text)
        writer.writerow(safe)
    return buffer.getvalue(), tuple(neutralized)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        raise HubError("mapping_rejected", "numeric cells must be text identifiers")
    raise HubError("mapping_rejected", "workbook cell type is not supported")


def _xlsx_digest(path: Path) -> str:
    sidecar = path.with_suffix(path.suffix + ".digest")
    if not sidecar.is_file():
        return ""
    return sidecar.read_text(encoding="utf-8").strip()
